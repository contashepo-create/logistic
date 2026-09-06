# -*- coding: utf-8 -*-
"""
اختبار طبقة الأمان وتعقيم المدخلات والتحقق من التواريخ والتفقيط والأرشيف.

هذه الطبقات أُضيفت بعد المراجعة الأمنية/المحاسبية لمطابقة نسخة الويب،
وهذا الملف يثبّتها ضد الانحدار.

تشغيل:
  python scripts/test_security.py
"""
from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
os.environ["LOGISTIC_DATA_DIR"] = tempfile.mkdtemp(prefix="logistic_sec_")

from app.core import calc, db, repo          # noqa: E402
from app.core.rules import RuleError          # noqa: E402
from app.core.security import (               # noqa: E402
    looks_malicious, neutralize_formula, sanitize_text, strip_control_chars,
)
from app.utils import fmt                     # noqa: E402

PASS = 0
FAILS: list[str] = []


def check(name: str, cond: bool, extra: str = "") -> None:
    global PASS
    if cond:
        PASS += 1
        print(f"✅ {name} {extra}")
        return
    FAILS.append(name)
    print(f"❌ {name} {extra}")


def reject(name: str, fn) -> None:
    try:
        fn()
    except (RuleError, ValueError):
        check(name, True)
        return
    except Exception as e:  # noqa: BLE001
        check(name, False, f"— استثناء غير متوقع {type(e).__name__}")
        return
    check(name, False, "— لم يُرفض!")


def main() -> None:
    db.init_db()
    conn = db.get_conn()

    # ------------------------------------------------ 1) التحقق من التواريخ
    print("== 1) التحقق من التواريخ (مطابقة safeIsoDate/safeFinancialYear)")
    reject("صيغة تاريخ خاطئة تُرفض",
           lambda: repo.save_year(conn, {"year": 2030, "date_from": "2030-01-01",
                                         "date_to": "31-12-2030"}))
    reject("تاريخ مستحيل (30 فبراير) يُرفض",
           lambda: repo.save_year(conn, {"year": 2030, "date_from": "2030-02-01",
                                         "date_to": "2030-02-30"}))
    reject("سنة أقصر من 180 يوماً تُرفض",
           lambda: repo.save_year(conn, {"year": 2030, "date_from": "2030-01-01",
                                         "date_to": "2030-01-04"}))
    reject("سنة أطول من 550 يوماً تُرفض",
           lambda: repo.save_year(conn, {"year": 2030, "date_from": "2030-01-01",
                                         "date_to": "2032-06-30"}))
    reject("سنة خارج 1900–2200 تُرفض",
           lambda: repo.save_year(conn, {"year": 1850, "date_from": "1850-01-01",
                                         "date_to": "1850-12-31"}))
    reject("رقم سنة لا يطابق تاريخ البداية يُرفض",
           lambda: repo.save_year(conn, {"year": 2099, "date_from": "2030-01-01",
                                         "date_to": "2030-12-31"}))
    y = repo.save_year(conn, {"year": 2030, "date_from": "2030-01-01",
                              "date_to": "2030-12-31"})
    check("سنة صالحة تُقبل", y > 0)

    # ------------------------------------------------ 2) الأنماط الهجومية
    print("== 2) رفض الأنماط الهجومية عند الإدخال")
    attacks = ["'; DROP TABLE customers; --", "<script>alert(1)</script>",
               "javascript:alert(1)", "onerror=evil", "data:text/html,x",
               "../../../../etc/passwd", "${jndi:ldap}", "{{7*7}}",
               "'; UPDATE customers SET x=1; --", "xp_cmdshell('x')"]
    for i, p in enumerate(attacks):
        reject(f"نمط هجومي #{i} ({p[:28]})",
               lambda p=p: repo.save_customer(conn, {"name": p,
                                                     "opening_balance": 0}))
    check("looks_malicious تكتشف كل الأنماط",
          all(looks_malicious(p) for p in attacks))

    benign = ["عميل عادي", "مؤسسة النقل الحديثة", "O'Brien", "شركة أ.ب.ج",
              "a' OR '1'='1", "—", "🌍🚛", "%s%s%s"]
    for i, p in enumerate(benign):
        cid = repo.save_customer(conn, {"name": p, "opening_balance": 0})
        got = repo.get_customer(conn, cid)["name"]
        check(f"نص سليم #{i} يُخزن كما هو", got == p, f"({p!r} → {got!r})")

    # ------------------------------------------------ 3) محارف التحكم
    print("== 3) نزع محارف التحكم والاتجاهات الخفية")
    check("نزع محارف التحكم", strip_control_chars("أ\u0000ب\u0007ج") == "أبج")
    check("نزع محارف الاتجاه (bidi override)",
          strip_control_chars("حساب\u202e4019\u202c") == "حساب4019")
    check("نزع BOM", strip_control_chars("\ufeffنص") == "نص")
    cid = repo.save_customer(conn, {"name": "عميل\u202eخفي\u0000",
                                    "opening_balance": 0})
    check("الاسم يُخزن منزوع المحارف الخفية",
          repo.get_customer(conn, cid)["name"] == "عميلخفي")

    # ------------------------------------------------ 4) تعقيم النص
    print("== 4) تعقيم النص (لا وسوم ولا كيانات مموّهة)")
    check("وسوم HTML تُزال", sanitize_text("<b>عادي</b>") == "عادي")
    check("كيانات HTML المموّهة تُزال", sanitize_text("a&#60;b&gt;c") == "a b c")
    check("الطول يُحترم", len(sanitize_text("أ" * 500, 50)) == 50)
    check("المسافات المضخّمة تُضغط", sanitize_text("أ     ب") == "أ  ب")

    # ------------------------------------------------ 5) حقن صيغ Excel
    print("== 5) تحييد حقن صيغ Excel")
    for bad in ("=cmd|' /C calc'!A0", "+1+1", "-2+3", "@SUM(A1)", "\t=x", "\r=y"):
        check(f"صيغة {bad[:16]!r} تُحيَّد",
              neutralize_formula(bad) == "'" + bad)
    for good in ("عميل عادي", "1500.00", "", "مؤسسة"):
        check(f"نص {good[:12]!r} لا يتغير", neutralize_formula(good) == good)
    check("الأرقام لا تُحيَّد", neutralize_formula(1500.0) == 1500.0)

    # تصدير فعلي: الخلية تُكتب مُحيَّدة
    os.environ["QT_QPA_PLATFORM"] = "minimal"
    from PySide6.QtWidgets import QApplication
    QApplication.instance() or QApplication([])
    import app.utils.exporter as ex
    out = Path(os.environ["LOGISTIC_DATA_DIR"]) / "sec.xlsx"
    ex._ask_save_path = lambda *a, **k: str(out)
    ex._ask_overwrite = lambda *a, **k: True
    import PySide6.QtWidgets as _qw
    _qw.QMessageBox.information = staticmethod(lambda *a, **k: None)
    ex.export_excel(None, conn, "أمان", ["ن", "م"],
                    [["=1+1", "1500.00"], ["@SUM(A1)", "عميل"]], "sec")
    from openpyxl import load_workbook
    ws = load_workbook(str(out)).active
    found = {ws.cell(row=r, column=1).value
             for r in range(1, ws.max_row + 1)}
    check("ملف Excel: الصيغ مُحيَّدة", "'=1+1" in found and "'@SUM(A1)" in found)
    check("ملف Excel: الأرقام تبقى رقمية لا نصية",
          any(isinstance(ws.cell(row=r, column=2).value, (int, float))
              for r in range(1, ws.max_row + 1)))

    # ------------------------------------------------ 6) المرفقات
    print("== 6) أمان المرفقات")
    base = Path(os.environ["LOGISTIC_DATA_DIR"])
    src = base / "a.txt"
    src.write_text("محتوى")
    rel = repo.store_attachment(str(src))
    check("مرفق سليم يُخزن داخل مجلد المرفقات",
          rel.startswith("attachments")
          and (db.attachments_dir() / "a.txt").exists())
    check("مرفق مكرر يأخذ اسماً فريداً",
          repo.store_attachment(str(src)) != rel)
    reject("مرفق بمسار اجتياز يُرفض",
           lambda: repo.store_attachment(str(base / "a.txt" / "..")))
    reject("مرفق غير موجود يُرفض",
           lambda: repo.store_attachment(str(base / "none.txt")))

    # ------------------------------------------------ 7) التفقيط
    print("== 7) التفقيط (المبلغ كتابةً)")
    cases = {
        0: "صفر ريال فقط لا غير",
        1: "واحد ريال فقط لا غير",
        2: "اثنان ريال فقط لا غير",
        10: "عشرة ريال فقط لا غير",
        11: "أحد عشر ريال فقط لا غير",
        20: "عشرون ريال فقط لا غير",
        100: "مائة ريال فقط لا غير",
        125: "مائة وخمسة وعشرون ريال فقط لا غير",
        1000: "ألف ريال فقط لا غير",
        2000: "ألفان ريال فقط لا غير",
        3500: "ثلاثة آلاف وخمسمائة ريال فقط لا غير",
    }
    for value, expected in cases.items():
        got = fmt.amount_to_arabic_words(value, "ريال")
        check(f"تفقيط {value}", got == expected, f"({got})")
    check("تفقيط بالكسور",
          fmt.amount_to_arabic_words(1234567.89, "ريال")
          == "مليون ومائتان وأربعة وثلاثون ألفاً وخمسمائة وسبعة وستون ريال "
             "و تسعة وثمانون هللة فقط لا غير".replace(" و ", " و"))
    check("تفقيط سالب", fmt.amount_to_arabic_words(-45.5, "ريال")
          .startswith("سالب "))
    check("جانب الرصيد", fmt.balance_side(1500) == "debit"
          and fmt.balance_side(-300) == "credit" and fmt.balance_side(0) == "zero")
    check("نص الرصيد", fmt.balance_text(1500).endswith("(عليه)")
          and fmt.balance_text(-300).endswith("(له)")
          and "مسدَّد" in fmt.balance_text(0))

    # ------------------------------------------------ 8) أرشيف السلف والخصومات
    print("== 8) أرشيف السلفيات والخصومات مع تفصيل التسويات")
    cb = repo.save_account(conn, "cashbox", {"name": "خزينة", "created_date":
                                             "2030-01-01",
                                             "opening_balance": 900000})
    emp = repo.save_employee(conn, {"name": "سائق", "emp_type": "driver",
                                    "base_salary": 4000})
    adv = repo.save_payment(conn, {"date": "2030-01-05", "account_kind": "cashbox",
                                   "account_id": cb, "voucher_type": "advance",
                                   "employee_id": emp, "amount": 1000})
    ded = repo.save_deduction(conn, {"date": "2030-01-07", "employee_id": emp,
                                     "amount": 300, "reason": "تلفيات"})
    repo.save_payroll(conn, {"date": "2030-01-31", "employee_id": emp,
                             "period_year": 2030, "period_month": 1,
                             "account_kind": "cashbox", "account_id": cb,
                             "base_salary": 4000, "settlements": [(adv, 600)],
                             "deduction_settlements": [(ded, 300)]})
    row = next(r for r in calc.advance_archive(conn, emp) if r["id"] == adv)
    check("أرشيف السلفة: المسدد والمتبقي",
          row["settled"] == 600 and row["remaining"] == 400
          and row["status"] == "partial")
    check("أرشيف السلفة: تفصيل المسيرة",
          len(row["settlements"]) == 1
          and row["settlements"][0]["amount"] == 600
          and row["settlements"][0]["period_label"] == "يناير 2030")
    tot = calc.advance_archive_totals(calc.advance_archive(conn, emp))
    check("إجماليات الأرشيف", tot == {"total": 1000.0, "settled": 600.0,
                                      "remaining": 400.0, "open_count": 1},
          str(tot))
    drow = calc.deduction_archive(conn, emp)[0]
    check("أرشيف الخصم: مسدَّد بالكامل",
          drow["status"] == "settled" and drow["remaining"] == 0
          and drow["settlements"][0]["amount"] == 300)

    print()
    if FAILS:
        print(f"💥 فشل {len(FAILS)} من {PASS + len(FAILS)}: {'، '.join(FAILS)}")
        sys.exit(1)
    print(f"🎉 كل اختبارات الأمان والسلامة نجحت ({PASS} فحصاً).")


if __name__ == "__main__":
    main()
